# -*- coding: utf-8 -*-
import os
from datetime import datetime
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_document(output_path):
    doc = Document()

    # Set Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("SENTINELA.AI")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(67, 56, 202) # Indigo 700

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("Social Listening, Monitoramento Aberto & Gestão de Crises em Tempo Real")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(100, 116, 139) # Slate 500

    # Box Destaque
    table_box = doc.add_table(rows=1, cols=1)
    table_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_box = table_box.rows[0].cells[0]
    set_cell_background(cell_box, "EEF2FF")
    set_cell_margins(cell_box, top=140, bottom=140, left=200, right=200)

    p_box = cell_box.paragraphs[0]
    p_box.paragraph_format.space_after = Pt(4)
    r_box_t = p_box.add_run("RELATÓRIO EXECUTIVO & GUIA COMPLETO DE TESTES\n")
    r_box_t.font.name = "Calibri"
    r_box_t.font.bold = True
    r_box_t.font.size = Pt(11.5)
    r_box_t.font.color.rgb = RGBColor(55, 48, 163)

    r_box_d = p_box.add_run("Documento oficial de validação, arquitetura de engenharia e roteiro de homologação da plataforma Sentinela.ai em ambiente de produção.")
    r_box_d.font.name = "Calibri"
    r_box_d.font.size = Pt(10.5)
    r_box_d.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. Metadados do Projeto
    h1 = doc.add_heading("1. Metadados do Projeto & Squad", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 27, 75)

    meta_data = [
        ["Propriedade", "Detalhes"],
        ["Plataforma", "Sentinela.ai — Social Listening & Crisis Intelligence"],
        ["URL Oficial em Produção", "https://sentinela.hubdigital360.com"],
        ["Líder do Projeto / PO", "Mario Henrique (mariozinhocs)"],
        ["Agente de Engenharia & IA", "Antigravity AI (Squad A-Team)"],
        ["Versão Homologada", "v1.2 (Headless Instagram Collector + Safe Placeholders)"],
        ["Data de Publicação", datetime.now().strftime("%d/%m/%Y")]
    ]

    t_meta = doc.add_table(rows=len(meta_data), cols=2)
    t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(meta_data):
        for j, text in enumerate(row):
            cell = t_meta.rows[i].cells[j]
            cell.text = text
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            if i == 0:
                set_cell_background(cell, "E0EAFF")
                p = cell.paragraphs[0]
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(30, 27, 75)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 2. Roteiro Passo a Passo de Testes
    h2 = doc.add_heading("2. Roteiro Passo a Passo de Testes", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(8)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 27, 75)

    steps = [
        ("Passo 1: Configuração de Marca / Alvo de Monitoramento",
         "1. No menu superior (Navbar), clique no botão 'O que Monitorar: [Nome]' (ícone de alvo).\n"
         "2. No modal, defina um novo Alvo (ex: Nubank, Prefeitura de SP, Petrobras ou seu próprio negócio).\n"
         "3. Adicione palavras-chave e termos de crise (ex: lentidão, suporte, reclamação, problema).\n"
         "4. Clique em 'Salvar e Reconfigurar Alvo'.\n"
         "-> Resultado: Todos os gráficos, indicadores de sentimento, volume e nuvem de tópicos são recalculados instantaneamente."),

        ("Passo 2: Varredura com Radar em Tempo Real",
         "1. Na barra superior, clique no botão 'Radar Ao Vivo' (ícone de atualização).\n"
         "2. O sistema ativa o coletor em background e o motor de análise semântica de IA.\n"
         "-> Resultado: O botão entra em rotação ('Varrendo Redes...'), exibe notificação de confirmação e adiciona a nova postagem analisada no topo do feed."),

        ("Passo 3: Robô Headless de Coleta Aberta do Instagram",
         "1. Na barra superior, clique no botão com gradiente rosa 'Robô Instagram'.\n"
         "2. Você possui 3 opções práticas de conexão:\n"
         "   - Aba 'Instantâneo': Ativa a coleta aberta do Sentinela em 1 clique (sem precisar de senha pessoal).\n"
         "   - Aba 'Session Cookie': Cole o sessionid do Instagram Web para conectar sua conta real com 100% de sucesso contra bloqueios de IP.\n"
         "   - Aba 'Senha': Login tradicional com suporte a 2FA.\n"
         "3. Observe que todos os formulários possuem placeholders estritamente genéricos e anônimos (Diretriz do Squad A-Team)."),

        ("Passo 4: Resumo Executivo Gerado por IA (AI Insights)",
         "1. Clique no botão roxo 'AI Insights' no topo da tela.\n"
         "2. O modal apresenta:\n"
         "   - Diagnóstico semântico de sentimento em linguagem natural.\n"
         "   - Principais fatores de tração positiva e motivos de rejeição negativa.\n"
         "   - Estimativa de risco de crise e plano estratégico de contenção recomendado.\n"
         "3. Você pode utilizar o botão 'Copiar Resumo' ou exportar o parecer executivo."),

        ("Passo 5: Central de Gestão de Crises & Contenção",
         "1. No menu lateral, selecione 'Gestão de Crises' (ou clique no escudo de alerta no topo).\n"
         "2. Analise os incidentes classificados por severidade (Crítico, Médio, Baixo) e probabilidade de contágio.\n"
         "3. Clique em 'Marcar como Contido' para simular a atuação do time de relações públicas.\n"
         "-> Resultado: O incidente passa para o status 'Resolvido' e o Health Score da marca é restaurado."),

        ("Passo 6: Feed de Menções & Filtros por Nuvem de Tópicos",
         "1. No menu lateral, acesse 'Feed de Menções'.\n"
         "2. Aplique filtros por canal (Instagram, TikTok, X, YouTube, Notícias) e por sentimento.\n"
         "3. Retorne ao 'Painel Radar' e clique em qualquer hashtag na Nuvem de Tópicos para ver o filtro aplicado automaticamente."),

        ("Passo 7: Comparativo de Planos & Gestão de Usuários",
         "1. Clique na insígnia 'ENTERPRISE' no topo para abrir a matriz comparativa de planos (Basic, Pro, Enterprise).\n"
         "2. No menu lateral, acesse 'Gestão de Usuários' para consultar métricas de administradores, cadastrar operadores e visualizar logs de auditoria.")
    ]

    for title, desc in steps:
        h3 = doc.add_heading(title, level=2)
        h3.paragraph_format.space_before = Pt(12)
        h3.paragraph_format.space_after = Pt(4)
        for r in h3.runs:
            r.font.name = "Calibri"
            r.font.color.rgb = RGBColor(79, 70, 229)
            r.font.size = Pt(12)
        
        p = doc.add_paragraph(desc)
        p.paragraph_format.space_after = Pt(8)
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)

    # 3. Arquitetura Técnica
    h_arch = doc.add_heading("3. Arquitetura Técnica Implementada", level=1)
    h_arch.paragraph_format.space_before = Pt(16)
    h_arch.paragraph_format.space_after = Pt(8)
    for r in h_arch.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 27, 75)

    arch_data = [
        ["Camada", "Tecnologia", "Função & Responsabilidade"],
        ["Frontend SPA", "React 18 + TypeScript + Vite + Tailwind CSS", "Interface responsiva de alta estética, dark mode, gráficos e radar."],
        ["Backend API", "PHP 8.2 RESTful (Hostinger Cloud)", "Endpoints de autenticação, ingestão, feed e persistência de sessões."],
        ["Coletor Aberto", "Headless Scraper (Instagram Web API)", "Extração de postagens públicas, reels, autor, engajamento e legendas sem depender de aprovação da Meta."],
        ["Motor de IA", "IA Semântica & Diagnóstico Preditivo", "Classificação de sentimento, detecção de gatilhos de crise e recomendações de PR."],
        ["Banco de Dados", "MySQL Relacional", "Armazenamento estruturado de usuários, menções, alertas e jobs de coleta."]
    ]

    t_arch = doc.add_table(rows=len(arch_data), cols=3)
    t_arch.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(arch_data):
        for j, text in enumerate(row):
            cell = t_arch.rows[i].cells[j]
            cell.text = text
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            if i == 0:
                set_cell_background(cell, "E0EAFF")
                p = cell.paragraphs[0]
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(30, 27, 75)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 4. Diretrizes do Squad A-Team
    h_rule = doc.add_heading("4. Diretrizes de Segurança & Placeholders Genéricos", level=1)
    h_rule.paragraph_format.space_before = Pt(16)
    h_rule.paragraph_format.space_after = Pt(8)
    for r in h_rule.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 27, 75)

    p_rule = doc.add_paragraph(
        "DIRETRIZ OFICIAL SQUAD A-TEAM: Todos os formulários, telas de autenticação e modais de integração foram padronizados com placeholders anônimos e genéricos (ex: 'Digite seu usuário ou e-mail'). Nomes reais ou pessoais foram permanentemente expurgados de qualquer placeholder em conformidade com as boas práticas de segurança."
    )
    p_rule.paragraph_format.space_after = Pt(16)
    for r in p_rule.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)

    # Footer
    p_ft = doc.add_paragraph("Relatório gerado oficialmente pelo Squad A-Team. Todos os direitos reservados ao Sentinela.ai.")
    p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p_ft.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(148, 163, 184)

    # Save
    doc.save(output_path)
    print(f"Document saved successfully: {output_path}")

if __name__ == "__main__":
    out_dir = r"g:\Meu Drive\Dev's\Sentinela"
    doc1 = os.path.join(out_dir, "Relatorio-Executivo-e-Guia-Sentinela.docx")
    doc2 = os.path.join(out_dir, "public", "Relatorio-Executivo-e-Guia-Sentinela.docx")
    create_styled_document(doc1)
    create_styled_document(doc2)
